"""Importing the owner's hand labels from the labelling workbook (BE-15, D-028)."""

from __future__ import annotations

import json

from openpyxl import load_workbook

from bookie_eval import labelling


def workbook(tmp_path, n: int = 4):
    books = [
        {"number": i + 1, "editionId": f"ed-{i}", "title": f"Book {i}", "authors": ["A"],
         "description": "d", "reviews": [{"number": 1, "rating": 4, "text": "fine"}]}
        for i in range(n)
    ]
    src = tmp_path / "books.json"
    src.write_text(json.dumps({"books": books}), "utf-8")
    path = tmp_path / "labels.xlsx"
    labelling.build(src, path)
    set_path = tmp_path / "set.json"
    set_path.write_text(json.dumps({"books": [{"editionId": b["editionId"]} for b in books]}), "utf-8")
    return path, set_path, [b["editionId"] for b in books]


def fill(path, rows: dict[int, list]):
    """Writes Pace, Tone, Content intensity, Theme 1 to Theme 5 (columns F to M) on Labels rows."""
    wb = load_workbook(path)
    ws = wb["Labels"]
    for r, values in rows.items():
        for c, v in enumerate(values, start=6):
            ws.cell(row=r, column=c, value=v)
    wb.save(path)


def test_themes_are_a_set_in_any_column_and_finished_rows_are_written(tmp_path):
    path, set_path, ids = workbook(tmp_path)
    fill(path, {
        2: ["fast", "tense", "moderate", "survival", None, "Power and politics ", "survival", None],
        3: ["slow", "melancholic", "low", None, None, None, None, "family"],
        4: ["moderate", None, None, "friendship"],  # partly filled: not imported
    })
    out = tmp_path / "hand-traits.json"
    result = labelling.import_labels(path, set_path, out)
    assert result["problems"] == [] and result["incomplete"] == [4]
    data = json.loads(out.read_text("utf-8"))
    assert data["labelled"] == 2 and data["of"] == 4
    assert data["books"][ids[0]] == {
        "pace": "fast", "tone": "tense", "contentIntensity": "moderate",
        "themes": ["survival", "power and politics"],
    }
    # Theme 5 alone is as good as Theme 1.
    assert data["books"][ids[1]]["themes"] == ["family"]
    assert ids[2] not in data["books"] and ids[3] not in data["books"]


def test_values_outside_the_vocabulary_stop_the_import(tmp_path):
    path, set_path, _ = workbook(tmp_path)
    fill(path, {2: ["quick", "tense", "low", "dragons"]})
    out = tmp_path / "hand-traits.json"
    result = labelling.import_labels(path, set_path, out)
    assert any('pace "quick"' in p for p in result["problems"])
    assert any('theme "dragons"' in p for p in result["problems"])
    assert not out.exists()


def test_books_outside_the_set_are_refused(tmp_path):
    path, set_path, ids = workbook(tmp_path)
    set_path.write_text(json.dumps({"books": [{"editionId": e} for e in ids[1:]]}), "utf-8")
    fill(path, {2: ["fast", "tense", "low"]})
    result = labelling.import_labels(path, set_path, tmp_path / "out.json")
    assert any("not in eval/llm/set.json" in p for p in result["problems"])
